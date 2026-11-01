import os
import sys
import json
from openpyxl.utils import get_column_letter

# Ensure we can import the config
sys.path.append(os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))
try:
    from config.excel_styles import get_theme
except ImportError:
    # Fallback if config is missing
    from openpyxl.styles import Font, Border, Side, PatternFill, Alignment
    class DummyTheme:
        @staticmethod
        def get_theme(name):
            return {
                "header_font": Font(name="Calibri", size=11, bold=True),
                "header_fill": PatternFill(fill_type="solid", fgColor="D9D9D9"),
                "body_font": Font(name="Calibri", size=11),
                "body_fill": PatternFill(fill_type=None),
                "alignment": Alignment(horizontal="center", vertical="center", wrap_text=True),
                "border": Border(left=Side(style="thin"), right=Side(style="thin"), top=Side(style="thin"), bottom=Side(style="thin"))
            }
    get_theme = DummyTheme.get_theme

def apply_styles(workbook_bundle, serialized_document, theme_name="DEFAULT_THEME"):
    """
    Phase 4.5 - Style Manager
    Applies professional formatting to the workbook without modifying data or merges.
    Implements Enterprise Improvements: Centralized Config, Traces, Progress Logs, and Protection.
    """
    print("Applying styles...")
    
    wb = workbook_bundle.get("workbook")
    worksheet_map = workbook_bundle.get("worksheet_map", {})
    
    # 1 & 2. Load Theme objects (cached outside loop)
    theme = get_theme(theme_name)
    align_center = theme["alignment"]
    border_all = theme["border"]
    font_header = theme["header_font"]
    fill_header = theme["header_fill"]
    font_body = theme["body_font"]
    fill_body = theme["body_fill"]
    
    stats = {
        "worksheets": 0,
        "header_cells": 0,
        "body_cells": 0,
        "styled_cells": 0
    }
    
    style_trace = []
    
    worksheets_data = serialized_document.get("workbook", {}).get("worksheets", [])
    
    for ws_data in worksheets_data:
        ws_name = ws_data.get("name")
        ws = worksheet_map.get(ws_name)
        
        if not ws:
            continue
            
        print(f"Worksheet: {ws_name}")
        stats["worksheets"] += 1
        
        ws_trace = {
            "worksheet": ws_name,
            "cells": []
        }
        
        # 8. Protected Cells configuration
        protected_coordinates = set() # Designed for future expansion
        
        rows = ws_data.get("rows", [])
        
        expected_cells = 0
        for row in rows:
            expected_cells += len(row.get("cells", []))
            
        current_styled = 0
        
        for row_data in rows:
            cells = row_data.get("cells", [])
            
            for cell_data in cells:
                # 0-based indexing from layout geometry
                r = cell_data.get("row", 0)
                c = cell_data.get("column", 0)
                rowspan = cell_data.get("rowspan", 1)
                colspan = cell_data.get("colspan", 1)
                
                start_row = r + 1
                start_column = c + 1
                end_row = start_row + rowspan - 1
                end_column = start_column + colspan - 1
                
                # 3. Header Classification logic
                is_header = False
                logical_type = cell_data.get("logical_type")
                if logical_type == "header":
                    is_header = True
                elif logical_type is None and r == 0:
                    is_header = True
                    
                # Iterate over the entire bounding box
                for r_idx in range(start_row, end_row + 1):
                    for c_idx in range(start_column, end_column + 1):
                        coord = (r_idx, c_idx)
                        if coord in protected_coordinates:
                            continue
                            
                        excel_cell = ws.cell(row=r_idx, column=c_idx)
                        coord_ref = f"{get_column_letter(c_idx)}{r_idx}"
                        
                        # Apply universal base styles
                        excel_cell.alignment = align_center
                        excel_cell.border = border_all
                        
                        applied_style = "BODY"
                        if is_header:
                            excel_cell.font = font_header
                            excel_cell.fill = fill_header
                            stats["header_cells"] += 1
                            applied_style = "HEADER"
                        else:
                            excel_cell.font = font_body
                            excel_cell.fill = fill_body
                            stats["body_cells"] += 1
                            
                        stats["styled_cells"] += 1
                        
                        # 4. Style Trace
                        ws_trace["cells"].append({
                            "coordinate": coord_ref,
                            "style": applied_style
                        })
                        
                current_styled += 1
                
                # 5. Progress Logging
                if current_styled % 50 == 0 or current_styled == expected_cells:
                    print(f"Styled {current_styled}/{expected_cells}")
                    
        # 6. Integrity Validation
        if current_styled != expected_cells:
            print(f"Warning: Style mismatch detected in {ws_name}. Expected: {expected_cells}, Styled regions: {current_styled}")
            
        print("Worksheet completed.")
        style_trace.append(ws_trace)
                    
    # Save Style Statistics
    try:
        base_dir = os.path.abspath(os.path.dirname(__file__))
        debug_folder = os.path.join(base_dir, "..", "debug")
        os.makedirs(debug_folder, exist_ok=True)
        
        with open(os.path.join(debug_folder, "style_report.json"), "w") as f:
            json.dump({"style_statistics": stats}, f, indent=4)
            
        with open(os.path.join(debug_folder, "style_trace.json"), "w") as f:
            json.dump(style_trace, f, indent=4)
            
    except Exception as e:
        print(f"Warning: Failed to write debug style files: {e}")
        
    return {
        "workbook": wb,
        "worksheet_map": worksheet_map,
        "statistics": stats
    }
