import os
import json
from openpyxl.utils import get_column_letter

def apply_merges(workbook_bundle, serialized_document):
    """
    Phase 4.4 - Merge Manager
    Applies all merged cell regions to the populated workbook using the 
    rowspan and colspan information generated during the Layout Analysis pipeline.
    Includes Enterprise Improvements: Conflict Detection, Tracing, and Analytics.
    """
    print("Applying merges...")
    
    wb = workbook_bundle.get("workbook")
    worksheet_map = workbook_bundle.get("worksheet_map", {})
    
    global_stats = {
        "worksheets": 0,
        "merged_regions": 0,
        "horizontal_merges": 0,
        "vertical_merges": 0,
        "complex_merges": 0
    }
    
    merge_trace = []
    merge_metadata = []
    
    worksheets_data = serialized_document.get("workbook", {}).get("worksheets", [])
    
    for ws_data in worksheets_data:
        ws_name = ws_data.get("name")
        ws = worksheet_map.get(ws_name)
        
        if not ws:
            print(f"Warning: Worksheet {ws_name} found in serialized document but missing in workbook bundle.")
            continue
            
        print(f"Worksheet: {ws_name}")
        global_stats["worksheets"] += 1
        
        ws_trace = {
            "worksheet": ws_name,
            "merges": []
        }
        
        ws_metadata = {
            "worksheet": ws_name,
            "total_merges": 0,
            "largest_merge": None,
            "largest_merge_area": 0,
            "merge_types": {
                "horizontal": 0,
                "vertical": 0,
                "complex": 0
            }
        }
        
        # 1. Merge Conflict Detection state
        occupied_cells = set()
        
        expected_merges = 0
        applied_merges = 0
        
        rows = ws_data.get("rows", [])
        
        # First pass: Count expected merges
        for row_data in rows:
            for cell in row_data.get("cells", []):
                if cell.get("rowspan", 1) > 1 or cell.get("colspan", 1) > 1:
                    expected_merges += 1
                    
        current_merge_idx = 0
        
        for row_data in rows:
            cells = row_data.get("cells", [])
            
            for cell in cells:
                rowspan = cell.get("rowspan", 1)
                colspan = cell.get("colspan", 1)
                
                # Skip standard 1x1 cells
                if rowspan == 1 and colspan == 1:
                    continue
                    
                current_merge_idx += 1
                
                # 2. Invalid Merge Detection
                if rowspan < 1 or colspan < 1:
                    print(f"Warning: Invalid merge span (rowspan={rowspan}, colspan={colspan}). Skipping.")
                    continue
                    
                start_row = cell.get("row", 0) + 1
                start_column = cell.get("column", 0) + 1
                end_row = start_row + rowspan - 1
                end_column = start_column + colspan - 1
                
                if end_row < start_row or end_column < start_column:
                    print(f"Warning: Invalid merge boundaries ({start_row},{start_column}) to ({end_row},{end_column}). Skipping.")
                    continue
                    
                # 1. Merge Conflict Detection Execution
                conflict = False
                for r in range(start_row, end_row + 1):
                    for c in range(start_column, end_column + 1):
                        if (r, c) in occupied_cells:
                            conflict = True
                            break
                    if conflict:
                        break
                        
                if conflict:
                    print(f"Warning: Attempting merge from R{start_row}C{start_column} to R{end_row}C{end_column} overlaps with existing merge. Skipping.")
                    continue
                    
                # Mark region as occupied
                for r in range(start_row, end_row + 1):
                    for c in range(start_column, end_column + 1):
                        occupied_cells.add((r, c))
                        
                # 7. Merge Classification
                m_type = "unknown"
                if rowspan == 1 and colspan > 1:
                    m_type = "horizontal"
                    ws_metadata["merge_types"]["horizontal"] += 1
                    global_stats["horizontal_merges"] += 1
                elif rowspan > 1 and colspan == 1:
                    m_type = "vertical"
                    ws_metadata["merge_types"]["vertical"] += 1
                    global_stats["vertical_merges"] += 1
                else:
                    m_type = "complex"
                    ws_metadata["merge_types"]["complex"] += 1
                    global_stats["complex_merges"] += 1
                    
                start_cell_ref = f"{get_column_letter(start_column)}{start_row}"
                end_cell_ref = f"{get_column_letter(end_column)}{end_row}"
                merge_ref = f"{start_cell_ref}:{end_cell_ref}"
                
                # Execute Merge
                try:
                    # 4. Progress Logging
                    print(f"Merge {current_merge_idx}/{expected_merges} ({merge_ref})")
                    
                    # 6. Preserve Top-Left Cell Value automatically by openpyxl
                    ws.merge_cells(
                        start_row=start_row, 
                        start_column=start_column, 
                        end_row=end_row, 
                        end_column=end_column
                    )
                    applied_merges += 1
                    global_stats["merged_regions"] += 1
                    ws_metadata["total_merges"] += 1
                    
                    area = rowspan * colspan
                    if area > ws_metadata["largest_merge_area"]:
                        ws_metadata["largest_merge_area"] = area
                        ws_metadata["largest_merge"] = merge_ref
                        
                    # 3. Merge Trace
                    ws_trace["merges"].append({
                        "start": start_cell_ref,
                        "end": end_cell_ref,
                        "rowspan": rowspan,
                        "colspan": colspan,
                        "type": m_type
                    })
                    
                except Exception as e:
                    print(f"Warning: Failed to merge region on {ws_name} {merge_ref}: {e}")
                    
        # 5. Integrity Validation
        if expected_merges != applied_merges:
            print(f"Warning: Merge mismatch detected on {ws_name}. Expected: {expected_merges}, Applied: {applied_merges}")
            
        print("Worksheet completed.")
        merge_trace.append(ws_trace)
        
        # Clean up transient area counter
        if "largest_merge_area" in ws_metadata:
            del ws_metadata["largest_merge_area"]
        merge_metadata.append(ws_metadata)
                    
    # Save the Merge Statistics and Traces
    try:
        base_dir = os.path.abspath(os.path.dirname(__file__))
        debug_folder = os.path.join(base_dir, "..", "debug")
        os.makedirs(debug_folder, exist_ok=True)
        
        with open(os.path.join(debug_folder, "merge_report.json"), "w") as f:
            json.dump({"merge_statistics": global_stats}, f, indent=4)
            
        with open(os.path.join(debug_folder, "merge_trace.json"), "w") as f:
            json.dump(merge_trace, f, indent=4)
            
        with open(os.path.join(debug_folder, "merge_metadata.json"), "w") as f:
            json.dump(merge_metadata, f, indent=4)
            
    except Exception as e:
        print(f"Warning: Failed to write debug files: {e}")
        
    return {
        "workbook": wb,
        "worksheet_map": worksheet_map,
        "statistics": global_stats
    }
