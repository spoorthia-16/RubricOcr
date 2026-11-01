import os
import json
import re
from datetime import datetime
from openpyxl import Workbook

def sanitize_sheet_name(name):
    """
    Sanitizes invalid Excel worksheet names by replacing forbidden characters
    and enforcing the 31-character limit.
    """
    # Excel forbidden characters: \ / ? * [ ] :
    sanitized = re.sub(r'[\\/?*\[\]:]', '_', name)
    # 31 character limit
    return sanitized[:31]

def build_workbook(serialized_document):
    """
    Phase 4.2 - Workbook Builder
    Creates the workbook and worksheet structure from the serialized document.
    Only manages the workbook object and worksheets; does not write any cell values.
    """
    print("Creating workbook...")
    wb = Workbook()
    
    # Remove the default sheet that openpyxl creates
    if "Sheet" in wb.sheetnames:
        default_sheet = wb["Sheet"]
        wb.remove(default_sheet)
        
    worksheet_map = {}
    
    metadata = {
        "generator": {
            "engine": "openpyxl",
            "version": "1.0",
            "generated_at": datetime.utcnow().isoformat() + "Z"
        },
        "worksheet_count": 0,
        "worksheets": []
    }
    
    worksheets = serialized_document.get("workbook", {}).get("worksheets", [])
    
    if not worksheets:
        raise ValueError("Workbook Validation Failed: At least one worksheet must exist.")
    
    seen_names = set()
    
    for idx, ws_data in enumerate(worksheets):
        raw_name = ws_data.get("name", f"Page_{idx + 1}")
        ws_name = sanitize_sheet_name(raw_name)
        
        # Ensure no duplicate worksheet names after sanitization
        original_ws_name = ws_name
        counter = 1
        while ws_name in seen_names:
            suffix = f"_{counter}"
            ws_name = original_ws_name[:31 - len(suffix)] + suffix
            counter += 1
            
        seen_names.add(ws_name)
        worksheet_id = f"ws_{str(idx + 1).zfill(3)}"
        
        # Create worksheet preserving the order of insertion
        ws = wb.create_sheet(title=ws_name)
        worksheet_map[ws_name] = ws
        
        # Calculate max rows and columns for metadata
        rows = ws_data.get("rows", [])
        row_count = len(rows)
        col_count = 0
        
        for row in rows:
            for cell in row.get("cells", []):
                col_end = cell.get("column", 0) + cell.get("colspan", 1) - 1
                if col_end >= col_count:
                    col_count = col_end + 1
                    
        # Validate row and column counts are non-negative
        if row_count < 0 or col_count < 0:
            raise ValueError(f"Workbook Validation Failed: Worksheet {ws_name} has invalid dimensions.")
            
        print(f"Creating worksheet: {ws_name}")
        print(f"Rows: {row_count}")
        print(f"Columns: {col_count}")
        
        metadata["worksheets"].append({
            "worksheet_id": worksheet_id,
            "name": ws_name,
            "rows": row_count,
            "columns": col_count
        })
        
    metadata["worksheet_count"] = len(metadata["worksheets"])
    
    print("Workbook initialized successfully.")
    
    # Save Workbook Metadata
    try:
        base_dir = os.path.abspath(os.path.dirname(__file__))
        debug_folder = os.path.join(base_dir, "..", "debug")
        os.makedirs(debug_folder, exist_ok=True)
        
        with open(os.path.join(debug_folder, "workbook_metadata.json"), "w") as f:
            json.dump(metadata, f, indent=4)
            
    except Exception as e:
        print(f"Warning: Failed to write debug workbook_metadata.json: {e}")
        
    return {
        "workbook": wb,
        "worksheet_map": worksheet_map
    }
