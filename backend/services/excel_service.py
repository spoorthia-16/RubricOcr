from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

def generate_excel(validated_json, output_path):
    """
    Creates an Excel file from the strictly validated JSON representation of the document.
    This is completely decoupled from any OCR or bounding box clustering logic!
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Extracted Document"
    
    grid_data = []
    
    if not validated_json or "tables" not in validated_json:
        wb.save(output_path)
        return output_path, grid_data
        
    current_row = 1
    
    # 1. Write global metadata at the top of the sheet
    metadata = validated_json.get("metadata", {})
    if metadata:
        for key, value in metadata.items():
            cell = ws.cell(row=current_row, column=1, value=f"{key}: {value}")
            cell.font = Font(bold=True)
            
            # Keep frontend synced
            grid_data.append([{"value": f"{key}: {value}", "colSpan": 1, "rowSpan": 1}])
            current_row += 1
            
        current_row += 1 # Add spacing
        grid_data.append([]) 
    
    # 2. Iterate through each table perfectly reconstructed by GPT
    for table in validated_json.get("tables", []):
        
        # Optional Table Title
        title = table.get("title")
        if title:
            cell = ws.cell(row=current_row, column=1, value=title)
            cell.font = Font(bold=True, size=14)
            grid_data.append([{"value": title, "colSpan": 1, "rowSpan": 1}])
            current_row += 1
            
        headers = table.get("headers", [])
        
        # Write Headers
        header_row_data = []
        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=current_row, column=col_idx, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="E0E0E0", end_color="E0E0E0", fill_type="solid")
            header_row_data.append({"value": header, "colSpan": 1, "rowSpan": 1})
            
        if header_row_data:
            grid_data.append(header_row_data)
            current_row += 1
            
        # Write Rows
        for row in table.get("rows", []):
            row_data = []
            cells_dict = row.get("cells", {})
            
            for col_idx, header in enumerate(headers, start=1):
                val = cells_dict.get(header, "")
                ws.cell(row=current_row, column=col_idx, value=val)
                row_data.append({"value": str(val), "colSpan": 1, "rowSpan": 1})
                
            grid_data.append(row_data)
            current_row += 1
            
        # Add visual spacing between multiple tables
        current_row += 1
        grid_data.append([])

    # Auto-adjust column widths for better aesthetics
    for column_cells in ws.columns:
        length = max(len(str(cell.value)) if cell.value else 0 for cell in column_cells)
        ws.column_dimensions[column_cells[0].column_letter].width = length + 2

    wb.save(output_path)
    return output_path, grid_data