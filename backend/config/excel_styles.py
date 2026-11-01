from openpyxl.styles import Font, Border, Side, PatternFill, Alignment

# Shared Components
CENTER_ALIGNMENT = Alignment(
    horizontal="center",
    vertical="center",
    wrap_text=True
)

THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin")
)

# Themes
THEMES = {
    "DEFAULT_THEME": {
        "header_font": Font(name="Calibri", size=11, bold=True, color="000000"),
        "header_fill": PatternFill(fill_type="solid", fgColor="D9D9D9"),
        "body_font": Font(name="Calibri", size=11, bold=False, color="000000"),
        "body_fill": PatternFill(fill_type=None),
        "alignment": CENTER_ALIGNMENT,
        "border": THIN_BORDER
    },
    "DARK_THEME": {
        "header_font": Font(name="Calibri", size=11, bold=True, color="FFFFFF"),
        "header_fill": PatternFill(fill_type="solid", fgColor="333333"),
        "body_font": Font(name="Calibri", size=11, bold=False, color="FFFFFF"),
        "body_fill": PatternFill(fill_type="solid", fgColor="555555"),
        "alignment": CENTER_ALIGNMENT,
        "border": THIN_BORDER
    }
}

def get_theme(theme_name="DEFAULT_THEME"):
    """
    Returns the loaded theme configuration for use in the style manager.
    """
    return THEMES.get(theme_name, THEMES["DEFAULT_THEME"])
